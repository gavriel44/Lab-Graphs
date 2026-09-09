"""Read measurement columns by name without silently discarding bad rows."""

import csv
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from .data import Dataset


def _from_rows(
    rows: Iterable[Sequence[Any]], *, x: str, y: str, sy: str, sx: str | None
) -> Dataset:
    iterator = iter(rows)
    header = next(iterator, None)
    if header is None:
        raise ValueError("The table is empty; a header row is required.")
    columns = {"x": x, "y": y, "sy": sy}
    if sx is not None:
        columns["sx"] = sx
    indexes = {}
    for field, column in columns.items():
        if header.count(column) != 1:
            raise ValueError(f"Expected exactly one column named {column!r} for {field}.")
        indexes[field] = header.index(column)

    values: dict[str, list[float]] = {field: [] for field in columns}
    for row_number, row in enumerate(iterator, start=2):
        # Completely empty rows are not observations. Partially empty rows are errors.
        if not row or all(value is None or value == "" for value in row):
            continue
        for field, index in indexes.items():
            try:
                value = float(row[index])
            except (TypeError, ValueError, IndexError) as exc:
                raise ValueError(
                    f"Row {row_number}, column {columns[field]!r}: expected a number. "
                    "For Excel formulas, recalculate and save the workbook first."
                ) from exc
            values[field].append(value)
    return Dataset(**values, x_label=x, y_label=y)


def read_csv(
    path: str | Path,
    *,
    x: str,
    y: str,
    sy: str,
    sx: str | None = None,
    delimiter: str = ",",
) -> Dataset:
    """Read a UTF-8 CSV with a header row, selecting columns by their exact names.

    Blank rows are skipped; incomplete or nonnumeric observations are rejected.
    Omit ``sx`` for exact x. Labels, including units, come from the selected headers.
    """
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        return _from_rows(csv.reader(stream, delimiter=delimiter), x=x, y=y, sx=sx, sy=sy)


def read_excel(
    path: str | Path,
    *,
    x: str,
    y: str,
    sy: str,
    sx: str | None = None,
    sheet: str | int = 0,
) -> Dataset:
    """Read an .xlsx sheet by name or zero-based index; requires the excel extra.

    The first row must be a header. Formula cells use Excel's saved values;
    openpyxl does not calculate formulas. Legacy .xls files are not supported.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError('Excel support requires: pip install "labgraphs[excel]"') from exc

    if isinstance(sheet, bool) or not isinstance(sheet, (str, int)):
        raise ValueError("sheet must be a name or a nonnegative integer index.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if isinstance(sheet, int):
            if not 0 <= sheet < len(workbook.sheetnames):
                raise ValueError(f"Sheet index {sheet} is out of range.")
            sheet = workbook.sheetnames[sheet]
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Unknown sheet {sheet!r}. Available: {workbook.sheetnames}.")
        return _from_rows(workbook[sheet].values, x=x, y=y, sy=sy, sx=sx)
    finally:
        workbook.close()
