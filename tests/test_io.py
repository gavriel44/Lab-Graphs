import builtins

import pytest
from numpy.testing import assert_allclose
from openpyxl import Workbook

from labgraphs import read_csv, read_excel


def test_csv_selects_named_columns_and_accepts_bom_and_blank_rows(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text(
        "\ufeffnote;Voltage (V);Current (A);sy;sx\na;2.1;1;0.2;0.1\n\nb;4;2;0.2;0.1\n",
        encoding="utf-8",
    )
    data = read_csv(path, x="Current (A)", y="Voltage (V)", sx="sx", sy="sy", delimiter=";")
    assert_allclose(data.x, [1, 2])
    assert_allclose(data.y, [2.1, 4])
    assert_allclose(data.sx, [0.1, 0.1])
    assert data.y_label == "Voltage (V)"


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("", "empty"),
        ("x,y,sy\n", "nonempty"),
        ("x,y,other\n1,2,3", "column named 'sy'"),
        ("x,x,y,sy\n1,1,2,3", "exactly one column"),
        ("x,y,sy\n1,,0.1", "Row 2, column 'y'"),
        ("x,y,sy\n1", "Row 2, column 'y'"),
        ("x,y,sy\n1,nope,0.1", "expected a number"),
        ("x,y,sy\n1,NaN,0.1", "finite"),
    ],
)
def test_csv_rejects_invalid_tables(tmp_path, content, message):
    path = tmp_path / "bad.csv"
    path.write_text(content)
    with pytest.raises(ValueError, match=message):
        read_csv(path, x="x", y="y", sy="sy")


@pytest.fixture
def workbook_path(tmp_path):
    book = Workbook()
    book.active.title = "Empty"
    sheet = book.create_sheet("Measurements")
    sheet.append(["x", "y", "sy"])
    sheet.append([1, 2.1, 0.1])
    sheet.append([2, 3.9, 0.1])
    path = tmp_path / "measurements.xlsx"
    book.save(path)
    book.close()
    return path


@pytest.mark.parametrize("sheet", [1, "Measurements"])
def test_excel_sheet_selection(workbook_path, sheet):
    data = read_excel(workbook_path, x="x", y="y", sy="sy", sheet=sheet)
    assert_allclose(data.y, [2.1, 3.9])
    assert data.sx is None


@pytest.mark.parametrize("sheet", [-1, 4, "Missing", None, True])
def test_excel_rejects_invalid_sheet(workbook_path, sheet):
    with pytest.raises(ValueError, match="[Ss]heet"):
        read_excel(workbook_path, x="x", y="y", sy="sy", sheet=sheet)


def test_excel_formula_without_cached_value_is_reported(tmp_path):
    book = Workbook()
    book.active.append(["x", "y", "sy"])
    book.active.append([1, "=2*A2", 0.1])
    path = tmp_path / "formula.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="recalculate and save"):
        read_excel(path, x="x", y="y", sy="sy")


def test_excel_extra_has_actionable_error(monkeypatch):
    original = builtins.__import__

    def without_excel(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("not installed")
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_excel)
    with pytest.raises(ImportError, match=r"labgraphs\[excel\]"):
        read_excel("unused.xlsx", x="x", y="y", sy="sy")
